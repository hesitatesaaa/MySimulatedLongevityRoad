"""Generate the player-facing material grade reference for version 0.2.0."""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "我的模拟长生路_材料分级目录_0.2.0.docx"
OUTPUT.parent.mkdir(exist_ok=True)

doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.68)
section.left_margin = Inches(0.72)
section.right_margin = Inches(0.72)

styles = doc.styles
for style_name in ("Normal", "Title", "Heading 1", "Heading 2"):
    style = styles[style_name]
    style.font.name = "Microsoft YaHei"
    style.font.color.rgb = RGBColor(0, 0, 0)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
styles["Normal"].font.size = Pt(9.5)
styles["Normal"].paragraph_format.space_after = Pt(5)
styles["Title"].font.size = Pt(18)
styles["Title"].font.bold = True
styles["Title"].paragraph_format.space_after = Pt(9)
styles["Heading 1"].font.size = Pt(12)
styles["Heading 1"].font.bold = True
styles["Heading 1"].paragraph_format.space_before = Pt(12)
styles["Heading 1"].paragraph_format.space_after = Pt(5)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_border(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for side in ("top", "left", "bottom", "right"):
        edge = OxmlElement(f"w:{side}")
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "4")
        edge.set(qn("w:color"), "D9D9D9")
        borders.append(edge)


def set_margins(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for side, val in (("top", 75), ("bottom", 75), ("left", 90), ("right", 90)):
        tag = OxmlElement(f"w:{side}")
        tag.set(qn("w:w"), str(val))
        tag.set(qn("w:type"), "dxa")
        margins.append(tag)
    tc_pr.append(margins)


def table(headers, rows, widths):
    result = doc.add_table(rows=1, cols=len(headers))
    result.autofit = False
    for i, width in enumerate(widths):
        result.columns[i].width = Inches(width)
    for i, header in enumerate(headers):
        result.rows[0].cells[i].text = header
    for row in rows:
        cells = result.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
    for ri, row in enumerate(result.rows):
        for ci, cell in enumerate(row.cells):
            cell.width = Inches(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_border(cell)
            set_margins(cell)
            if ri == 0:
                shade(cell, "243D4C")
            elif ri % 2 == 0:
                shade(cell, "F1F5F7")
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.12
                if ci in (0, 2, 3) and len(headers) > 3:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.name = "Microsoft YaHei"
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
                    run.font.size = Pt(8.5)
                    if ri == 0:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
    header_tr_pr = result.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    header_tr_pr.append(repeat)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return result


doc.add_paragraph("我的模拟长生路 材料分级目录", "Title")
doc.add_paragraph("适用版本 0.2.0｜按当前物品目录与材料发现逻辑整理｜2026 年 9 月 25 日")
doc.add_paragraph(
    "当前十种炼丹与制符材料按稀有度分为黄、玄、地、天四阶。材料阶位决定发现范围和稀有度；丹药、符箓、法宝的制作品阶另行计算。材料在物品目录中的 Grade 字段均为 0，不能据此判定材料没有分级。"
)

doc.add_paragraph("一 材料分级总表", "Heading 1")
table(
    ["阶位", "ID", "材料", "价格", "偏好地形", "用途或说明"],
    [
        ("黄阶", "A06", "长生药材", 6, "林地", "长生丹药材的游戏化合并项"),
        ("黄阶", "A07", "聚灵髓", 6, "山地", "凝聚灵力，辅助突破"),
        ("黄阶", "A08", "护脉草", 6, "林地", "保护经脉，稳定伤势"),
        ("黄阶", "F01", "灵符纸", 6, "无", "每张符箓固定消耗一份"),
        ("玄阶", "A01", "琉璃珠", 12, "山地", "启灵、悟性及突破辅助"),
        ("玄阶", "A02", "蓝血珊瑚", 12, "水域", "冰寒、镇定及净体药性"),
        ("玄阶", "A03", "灵虚草", 12, "林地", "神魂恢复类药材"),
        ("玄阶", "A09", "洗髓液", 12, "山地", "洗髓及根基重塑"),
        ("地阶", "A04", "长生青力", 24, "无", "强生机、修复及重塑"),
        ("天阶", "A05", "万灵琼浆", 24, "无", "高阶突破和恢复稀有材料"),
    ],
    [0.66, 0.58, 1.04, 0.54, 0.89, 3.35],
)
doc.add_paragraph("价格单位为天玄镜中的贡献度。材料价格按物品目录记录；稀有度不单独决定价格。")

doc.add_paragraph("二 各材料的可发现来源", "Heading 1")
table(
    ["材料范围", "年度活动", "遗迹", "境界突破", "天地机缘", "势力委托"],
    [
        ("A01–A03、A06–A09", "可", "可", "可", "可", "可"),
        ("A04 长生青力", "—", "可", "可", "可", "可"),
        ("A05 万灵琼浆", "—", "可", "可", "可", "—"),
        ("F01 灵符纸", "可", "可", "—", "—", "可"),
    ],
    [1.75, 0.95, 0.71, 1.13, 1.13, 1.09],
)
doc.add_paragraph("“可”表示进入该来源的候选池，仍需满足事件触发、阶位上限与随机判定；“—”表示该来源不会产出。")

doc.add_paragraph("三 发现条件和概率", "Heading 1")
table(
    ["来源", "触发及阶位限制", "当前概率"],
    [
        ("年度活动", "符合修行资格的人物按年判定；仅黄、玄阶。偏好地形使对应材料概率翻倍。", "每种黄阶 0.30%，玄阶 0.08%；地形匹配时为 0.60% 和 0.16%。"),
        ("遗迹探索", "探索者存活且完成遗迹探索；地阶要求遗迹品质至少 3，天阶至少 4。", "每种黄阶 3%，玄阶 1%，地阶 0.4%，天阶 0.1%。"),
        ("境界突破", "真实向上突破；达金丹可抽到地阶，达元婴可抽到天阶。", "事件先以 2% 通过，再在合格材料中抽取一种。"),
        ("天地机缘", "成功获得洞天或天地之变奖励；品质至少 3 可抽地阶，至少 4 可抽天阶。", "事件先以 8% 通过，再在合格材料中抽取一种。"),
        ("势力委托", "完成资源奖励委托；最高品质委托另有地阶判定。", "黄或玄阶抽取门槛 3%；最高品质时另有地阶抽取门槛 0.5%。"),
    ],
    [1.0, 3.06, 2.99],
)
doc.add_paragraph(
    "突破、机缘和委托在通过门槛后，按候选材料的阶位权重抽取：黄 60、玄 30、地 9、天 1。该权重是每个候选材料的相对权重，不等于最终掉落概率。发现的材料进入人物乾坤袋，并记入天地资源事件。"
)
doc.add_paragraph(
    "资料口径：code/MySimulatedLongevityRoad/Data/MclslItemCatalog.cs；code/MySimulatedLongevityRoad/Systems/Crafting/MclslMaterialDiscovery.cs。"
)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run("我的模拟长生路 0.2.0  材料分级目录")
footer.runs[0].font.size = Pt(8)
footer.runs[0].font.color.rgb = RGBColor(100, 100, 100)

doc.save(OUTPUT)
print(OUTPUT)
