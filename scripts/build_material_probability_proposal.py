"""Create a review-only proposal for material acquisition increases."""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "我的模拟长生路_材料获取概率上调提案_第二版_待确认.docx"
OUT.parent.mkdir(exist_ok=True)

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(0.67)
sec.left_margin = sec.right_margin = Inches(0.70)

for name in ("Normal", "Title", "Heading 1"):
    sty = doc.styles[name]
    sty.font.name = "Microsoft YaHei"
    sty._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    sty.font.color.rgb = RGBColor(0, 0, 0)
doc.styles["Normal"].font.size = Pt(9)
doc.styles["Normal"].paragraph_format.space_after = Pt(4)
doc.styles["Title"].font.size = Pt(17)
doc.styles["Title"].font.bold = True
doc.styles["Title"].paragraph_format.space_after = Pt(7)
doc.styles["Heading 1"].font.size = Pt(11.5)
doc.styles["Heading 1"].font.bold = True
doc.styles["Heading 1"].paragraph_format.space_before = Pt(10)
doc.styles["Heading 1"].paragraph_format.space_after = Pt(4)


def cell_format(cell, header, alt):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for side, size in (("top", 65), ("bottom", 65), ("left", 90), ("right", 90)):
        edge = OxmlElement("w:" + side)
        edge.set(qn("w:w"), str(size))
        edge.set(qn("w:type"), "dxa")
        margins.append(edge)
    tc_pr.append(margins)
    borders = OxmlElement("w:tcBorders")
    for side in ("top", "bottom", "left", "right"):
        edge = OxmlElement("w:" + side)
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "4")
        edge.set(qn("w:color"), "D9D9D9")
        borders.append(edge)
    tc_pr.append(borders)
    if header or alt:
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "243D4C" if header else "F1F5F7")
        tc_pr.append(shd)
    for p in cell.paragraphs:
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.12
        for run in p.runs:
            run.font.name = "Microsoft YaHei"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
            run.font.size = Pt(8.3)
            if header:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)


def add_table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    for i, width in enumerate(widths):
        t.columns[i].width = Inches(width)
    for i, value in enumerate(headers):
        t.rows[0].cells[i].text = value
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
    for ri, row in enumerate(t.rows):
        for ci, cell in enumerate(row.cells):
            cell.width = Inches(widths[ci])
            cell_format(cell, ri == 0, ri % 2 == 0 and ri > 0)
    tr_pr = t.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def pct(value, digits=3):
    return f"{value:.{digits}f}".rstrip("0").rstrip(".") + "%"


doc.add_paragraph("我的模拟长生路 材料获取概率上调提案", "Title")
doc.add_paragraph("第二版｜0.2.0 当前代码对照｜待确认方案｜2026 年 9 月 25 日")
doc.add_paragraph(
    "建议将五类材料来源的获取机会提高到更容易感知的水平，并提高抽取池中玄、地、天阶材料的相对权重。十种材料的阶位、来源资格、地形加成、境界与地点品质门槛维持现状。本文件仅供确认；当前游戏代码尚未应用新概率。"
)

doc.add_paragraph("一 拟调整的判定值", "Heading 1")
add_table(
    ["来源及判定", "当前", "建议", "相对增幅"],
    [
        ("年度活动 每种黄阶", "0.30%", "1.00%", "+233%"),
        ("年度活动 每种玄阶", "0.08%", "0.30%", "+275%"),
        ("遗迹 每种黄阶", "3.00%", "8.00%", "+167%"),
        ("遗迹 每种玄阶", "1.00%", "3.00%", "+200%"),
        ("遗迹 每种地阶", "0.40%", "1.50%", "+275%"),
        ("遗迹 每种天阶", "0.10%", "0.50%", "+400%"),
        ("真实突破 材料抽取门槛", "2.00%", "6.00%", "+200%"),
        ("天地机缘 材料抽取门槛", "8.00%", "18.00%", "+125%"),
        ("势力委托 黄玄阶抽取门槛", "3.00%", "8.00%", "+167%"),
        ("最高品质委托 地阶抽取门槛", "0.50%", "2.00%", "+300%"),
    ],
    [3.35, 1.18, 1.18, 1.31],
)
doc.add_paragraph(
    "年度活动在材料偏好地形中仍翻倍：黄阶由 0.60% 调至 2.00%，玄阶由 0.16% 调至 0.60%。遗迹地阶仍需品质至少 3，天阶至少 4。"
)

doc.add_paragraph("二 每种材料的直接判定概率", "Heading 1")
add_table(
    ["阶位与材料", "年度活动 普通地形", "遗迹探索"],
    [
        ("黄 A06 长生药材", "0.30% → 1.00%", "3.00% → 8.00%"),
        ("黄 A07 聚灵髓", "0.30% → 1.00%", "3.00% → 8.00%"),
        ("黄 A08 护脉草", "0.30% → 1.00%", "3.00% → 8.00%"),
        ("黄 F01 灵符纸", "0.30% → 1.00%", "3.00% → 8.00%"),
        ("玄 A01 琉璃珠", "0.08% → 0.30%", "1.00% → 3.00%"),
        ("玄 A02 蓝血珊瑚", "0.08% → 0.30%", "1.00% → 3.00%"),
        ("玄 A03 灵虚草", "0.08% → 0.30%", "1.00% → 3.00%"),
        ("玄 A09 洗髓液", "0.08% → 0.30%", "1.00% → 3.00%"),
        ("地 A04 长生青力", "不可获得", "0.40% → 1.50%"),
        ("天 A05 万灵琼浆", "不可获得", "0.10% → 0.50%"),
    ],
    [2.80, 2.11, 2.11],
)
doc.add_paragraph(
    "年度活动是每个合格人物每年对每种材料分别判定；遗迹是在一次成功探索中对每种合格材料分别判定。表内遗迹概率以达到相应品质门槛为前提。灵符纸无偏好地形，其年度概率不会翻倍。"
)

doc.add_paragraph("三 抽取型来源的每种材料实际概率", "Heading 1")
doc.add_paragraph(
    "以下按最高解锁条件计算：真实突破到元婴及以上、天地机缘品质至少 4、最高品质势力委托。数值是“通过门槛并选中指定材料”的单次事件概率；较低境界或品质会改变候选池。"
)
old_weights = {"黄": 60, "玄": 30, "地": 9, "天": 1}
new_weights = {"黄": 60, "玄": 35, "地": 18, "天": 6}
old_total = 3 * old_weights["黄"] + 4 * old_weights["玄"] + old_weights["地"] + old_weights["天"]
new_total = 3 * new_weights["黄"] + 4 * new_weights["玄"] + new_weights["地"] + new_weights["天"]
old_faction_total = 4 * old_weights["黄"] + 4 * old_weights["玄"]
new_faction_total = 4 * new_weights["黄"] + 4 * new_weights["玄"]
rows = []
for tier, names in (
    ("黄", "A06／A07／A08"),
    ("黄", "F01 灵符纸"),
    ("玄", "A01／A02／A03／A09"),
    ("地", "A04"),
    ("天", "A05"),
):
    old_w, new_w = old_weights[tier], new_weights[tier]
    br_old, br_new = 2 * old_w / old_total, 6 * new_w / new_total
    op_old, op_new = 8 * old_w / old_total, 18 * new_w / new_total
    if tier in ("黄", "玄"):
        fac_old, fac_new = 3 * old_w / old_faction_total, 8 * new_w / new_faction_total
        faction = f"{pct(fac_old)} → {pct(fac_new)}"
    elif tier == "地":
        faction = "0.50% → 2.00%"
    else:
        faction = "不可获得"
    breakthrough = "不可获得" if names == "F01 灵符纸" else f"{pct(br_old)} → {pct(br_new)}"
    opportunity = "不可获得" if names == "F01 灵符纸" else f"{pct(op_old)} → {pct(op_new)}"
    rows.append((f"{tier} {names}", breakthrough, opportunity, faction))
add_table(
    ["同阶每种材料", "境界突破", "天地机缘", "势力委托"],
    rows,
    [2.05, 1.68, 1.68, 1.61],
)
doc.add_paragraph(
    "抽取权重建议从每种材料的黄 60、玄 30、地 9、天 1 调至黄 60、玄 35、地 18、天 6。突破和机缘不抽灵符纸，最高阶候选池总权重从 310 变为 344；委托黄玄阶池包含灵符纸，总权重从 360 变为 380。最高品质委托的地阶判定仅对应 A04。表中概率四舍五入至 0.001%，不应与代码门槛值混用。"
)

doc.add_paragraph("四 确认后实施的范围", "Heading 1")
doc.add_paragraph(
    "确认本方案后，调整 MclslMaterialDiscovery.cs 中十处判定值、四个抽取权重及对应注释。材料目录、配方、价格、来源资格、去重规则和存档结构均不需要改动。"
)
doc.add_paragraph(
    "依据：code/MySimulatedLongevityRoad/Systems/Crafting/MclslMaterialDiscovery.cs 与 code/MySimulatedLongevityRoad/Data/MclslItemCatalog.cs。"
)

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run("0.2.0 材料获取概率提案  待确认")
run.font.size = Pt(8)
run.font.color.rgb = RGBColor(100, 100, 100)

doc.save(OUT)
print(OUT)
